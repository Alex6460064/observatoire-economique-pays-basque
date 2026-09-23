import zipfile

import duckdb
import pytest
from openpyxl import Workbook
from socle_territorial.bmo import lire_bassins, lire_bmo
from socle_territorial.geo import normaliser_nom, simplifier
from socle_territorial.journal import Journal
from socle_territorial.stockage import ecrire_parquet


@pytest.mark.parametrize(
    ("brut", "attendu"),
    [
        ("Saint-Jean-de-Luz", "SAINT JEAN DE LUZ"),
        ("ST JEAN DE LUZ", "SAINT JEAN DE LUZ"),
        ("SaintJean de Luz", "SAINTJEAN DE LUZ"),  # faute de saisie : pas de magie
        ("Ossès", "OSSES"),
        ("L'Hôpital-Saint-Blaise", "HOPITAL SAINT BLAISE"),
        ("Bayonne Cedex", "BAYONNE"),
        ("Ste-Engrâce", "SAINTE ENGRACE"),
        (None, ""),
    ],
)
def test_normaliser_nom(brut, attendu):
    assert normaliser_nom(brut) == attendu


def _classeur_bmo(path, entete, lignes, feuille="BMO_2026_open_data"):
    wb = Workbook()
    wb.active.title = "Description_des_variables"
    ws = wb.create_sheet(feuille)
    ws.append(entete)
    for r in lignes:
        ws.append(r)
    wb.save(path)


ENTETE_2026 = [
    "annee",
    "Code métier BMO",
    "Nom métier BMO",
    "Famille_met",
    "Lbl_fam_met",
    "REG",
    "NOM_REG",
    "Dept",
    "NomDept",
    "BE26",
    "NOMBE26",
    "clpe",
    "met",
    "xmet",
    "smet",
]
# Format 2023 : colonnes dans un autre ordre, suffixe BE23.
ENTETE_2023 = [
    "annee",
    "Code métier BMO",
    "Nom métier BMO",
    "Famille_met",
    "Lbl_fam_met",
    "BE23",
    "NOMBE23",
    "Dept",
    "NomDept",
    "REG",
    "NOM_REG",
    "met",
    "xmet",
    "smet",
]


def test_lire_bmo_filtre_et_secret(tmp_path):
    f = tmp_path / "bmo.xlsx"
    _classeur_bmo(
        f,
        ENTETE_2026,
        [
            [
                2026,
                "A0X40",
                "Agriculteurs",
                "Z",
                "Autres",
                "75",
                "NA",
                "64",
                "PA",
                7539,
                "PAYS BASQUE",
                "c",
                "120",
                "40",
                "*",
            ],
            [2026, "A0X40", "Agriculteurs", "Z", "Autres", "75", "NA", "64", "PA", 7538, "BEARN", "c", "*", "*", "*"],
            [2026, "A0X40", "Agriculteurs", "Z", "Autres", "75", "NA", "40", "Landes", 7201, "X", "c", "9", "1", "1"],
        ],
    )
    lignes = lire_bmo(f, departements={"64"})
    assert len(lignes) == 2
    assert lignes[0] == {
        "annee": 2026,
        "code_metier": "A0X40",
        "libelle_metier": "Agriculteurs",
        "code_famille": "Z",
        "libelle_famille": "Autres",
        "code_departement": "64",
        "code_bassin": "7539",
        "libelle_bassin": "PAYS BASQUE",
        "projets": 120,
        "projets_difficiles": 40,
        "projets_saisonniers": None,
        "projets_secret": False,
    }
    assert lignes[1]["projets"] is None and lignes[1]["projets_secret"] is True


def test_lire_bmo_format_2023_dans_un_zip(tmp_path):
    x = tmp_path / "Base_open_data_23.xlsx"
    _classeur_bmo(
        x,
        ENTETE_2023,
        [
            [
                2023,
                "A0Z40",
                "Agriculteurs salariés",
                "Z",
                "Autres",
                7539,
                "PAYS BASQUE",
                "64",
                "PA",
                "75",
                "NA",
                "10",
                "2",
                "3",
            ],
        ],
        feuille="BMO_2023_open_data",
    )
    z = tmp_path / "bmo2023.zip"
    with zipfile.ZipFile(z, "w") as zf:
        zf.write(x, "Donnees/Base_open_data_23.xlsx")
    (ligne,) = lire_bmo(z)
    assert ligne["code_bassin"] == "7539" and ligne["projets_saisonniers"] == 3


def test_lire_bmo_colonne_manquante(tmp_path):
    f = tmp_path / "bmo.xlsx"
    _classeur_bmo(f, ENTETE_2026[:-1], [])
    with pytest.raises(ValueError, match="smet"):
        lire_bmo(f)


def test_lire_bassins(tmp_path):
    f = tmp_path / "b.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["region", "dep", "code_commune", "lib_commune", "code_bassin_2026", "lib_bassin_2026"])
    ws.append([75, "64", "64102", "BAYONNE", 7539, "PAYS BASQUE"])
    wb.save(f)
    assert lire_bassins(f) == [
        {
            "code_commune": "64102",
            "libelle_commune": "BAYONNE",
            "code_bassin": "7539",
            "libelle_bassin": "PAYS BASQUE",
            "millesime_zonage": 2026,
        }
    ]


def test_ecrire_parquet_schema_stable_meme_vide(tmp_path):
    dest = tmp_path / "vide.parquet"
    assert ecrire_parquet([], dest, {"a": "VARCHAR", "n": "INTEGER"}) == 0
    cols = duckdb.sql(f"describe select * from '{dest.as_posix()}'").fetchall()
    assert [(c[0], c[1]) for c in cols] == [("a", "VARCHAR"), ("n", "INTEGER")]


def test_ecrire_parquet_types_explicites(tmp_path):
    dest = tmp_path / "x.parquet"
    ecrire_parquet([{"a": "01", "n": 3, "ignore": 1}], dest, {"a": "VARCHAR", "n": "INTEGER"})
    assert duckdb.sql(f"select * from '{dest.as_posix()}'").fetchall() == [("01", 3)]


def test_journal_trace_les_echecs(tmp_path):
    j = Journal(tmp_path / "journal.jsonl")
    with j.extraction("src", millesime="2026") as e:
        e.lignes = 10
    with pytest.raises(RuntimeError), j.extraction("src"):
        raise RuntimeError("boum")
    lignes = j.lire()
    assert [x["statut"] for x in lignes] == ["ok", "echec"]
    assert "boum" in lignes[1]["erreur"]
    assert j.derniere_reussie("src")["lignes"] == 10


def test_simplifier_retire_les_points_alignes_et_arrondit():
    carre = [[0, 0], [0.5, 0.00001], [1, 0], [1, 1], [0, 1], [0, 0]]
    fc = {
        "type": "FeatureCollection",
        "features": [
            {"type": "Feature", "properties": {"code": "x"}, "geometry": {"type": "Polygon", "coordinates": [carre]}}
        ],
    }
    (f,) = simplifier(fc, tolerance=0.001, decimales=2)["features"]
    assert f["geometry"]["coordinates"][0] == [[0, 0], [1, 0], [1, 1], [0, 1], [0, 0]]
    assert f["properties"] == {"code": "x"}


def test_simplifier_garde_un_anneau_valide():
    petit = [[0, 0], [1e-6, 0], [1e-6, 1e-6], [0, 0]]
    fc = {
        "type": "FeatureCollection",
        "features": [
            {"type": "Feature", "properties": {}, "geometry": {"type": "MultiPolygon", "coordinates": [[petit]]}}
        ],
    }
    (f,) = simplifier(fc, tolerance=1)["features"]
    assert len(f["geometry"]["coordinates"][0][0]) >= 4
